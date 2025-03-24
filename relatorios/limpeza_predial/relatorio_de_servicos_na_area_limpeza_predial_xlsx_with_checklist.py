import openpyxl
from servicos.headers_report_limpeza_predial import headers_report_services
from servicos.utils_limpeza_predial import colect_dados_fato_servico_limpeza_predial
from checklists.models import CheckListLimpezaPredial  # Import do modelo de checklist
from django.http import HttpResponse
from utils.utils import generate_id_random
from datetime import datetime
from django.shortcuts import redirect
from permissionscontrol.utils import verify_login

def exportar_relatorio_de_serivos_na_area_limpeza_predial_excel_with_checklist(request, userid, id_random, DataDeInicio,
                                                                              DataDeConclusao, Areas, TipoServico,
                                                                              ServicosEscalados, ColaboradoresEscalados, type):
    block = verify_login(request=request, userid=userid)

    if block:
        return redirect('logout')

    # Converte as datas de início e conclusão, tratando o caso onde a data não é fornecida
    DataDeInicio = datetime.strptime(DataDeInicio, '%Y-%m-%dT%H:%M') if DataDeInicio and DataDeInicio != "None" else None
    DataDeConclusao = datetime.strptime(DataDeConclusao, '%Y-%m-%dT%H:%M') if DataDeConclusao and DataDeConclusao != "None" else None

    # Converte os campos de Serviços e Colaboradores para listas
    ServicosEscalados = ServicosEscalados.split(',')
    ColaboradoresEscalados = ColaboradoresEscalados.split(',')

    # Coleta os dados
    dados = colect_dados_fato_servico_limpeza_predial(
        request=request,
        userid=userid,
        DataDeInicio=DataDeInicio,
        DataDeConclusao=DataDeConclusao,
        TipoServico=TipoServico,
        Areas=Areas,
        ServicosEscalados=ServicosEscalados,
        ColaboradoresEscalados=ColaboradoresEscalados,
        status=['Concluido']
    )

    # Configuração inicial do workbook
    wb = openpyxl.Workbook()
    ws_services = wb.active
    ws_services.title = "Relatório de Execução Limpeza Predial"
    ws_checklists = wb.create_sheet(title="Dados do Checklist")  # Cria a segunda aba

    # Cabeçalhos da aba de serviços
    headers = headers_report_services
    for col_num, header_title in enumerate(headers, 1):
        cell = ws_services.cell(row=1, column=col_num)
        cell.value = header_title

    # Aplica filtros baseados no tipo
    if type == 'catalogo_de_servicos':
        dados = dados.filter(id_random_servico=id_random)
    elif type == 'configuracao':
        dados = dados.filter(id_random_configuracao=id_random)
    elif type == 'areas':
        dados = dados.filter(id_random_area=id_random)

    # Preenche os dados na aba de serviços
    for row_num, row in enumerate(dados, start=2):
        row_data = [
            row.tipodeempresa,
            row.empresaprestadora,
            row.id_servico,
            row.tipo_agendamento,
            row.descricao_do_servico,
            row.servicos_solicitados,
            row.data_de_inicio.replace(tzinfo=None) if row.data_de_inicio else None,
            row.data_de_conclusao.replace(tzinfo=None) if row.data_de_conclusao else None,
            row.status_servico,
            row.area_atendida,
            row.area_total,
            row.localidade,
            row.unidade,
            row.id_servico,
            row.tempo_na_area,
            row.colaborador_envolvido,
            row.colaborador_envolvido_id_random,
            row.foto_conclusao,
            row.data_hora_chegada.replace(tzinfo=None) if row.data_hora_chegada else None,
            row.data_hora_retorno.replace(tzinfo=None) if row.data_hora_retorno else None,
        ]
        for col_num, value in enumerate(row_data, start=1):
            cell = ws_services.cell(row=row_num, column=col_num)
            cell.value = value

    # Define cabeçalhos para a aba de checklists
    checklist_headers = [
        "ID Agendamento", "Descrição do Serviço", "Descrição do Checklist",
        "Foto", "Atualizado em"
    ]
    for col_num, header_title in enumerate(checklist_headers, 1):
        cell = ws_checklists.cell(row=1, column=col_num)
        cell.value = header_title

    # Adiciona os dados de checklists à aba "Dados do Checklist"
    row_checklist = 2
    for dado in dados:
        dados_checklist = CheckListLimpezaPredial.objects.filter(servico_agendado__id_random=dado.id_random_agendamento)
        for checklist in dados_checklist:
            ws_checklists.cell(row=row_checklist, column=1, value=dado.id_agendamento)
            ws_checklists.cell(row=row_checklist, column=2, value=dado.descricao_do_servico)
            ws_checklists.cell(row=row_checklist, column=3, value=checklist.descricao)
            ws_checklists.cell(row=row_checklist, column=4, value=checklist.foto_comprovacao.url if checklist.foto_comprovacao else "N/A")
            ws_checklists.cell(row=row_checklist, column=5, value=checklist.atualizado_em.strftime("%d/%m/%Y %H:%M") if checklist.atualizado_em else "N/A")
            row_checklist += 1

    # Ajusta a largura das colunas para melhor legibilidade
    for col in range(1, len(headers) + 1):
        ws_services.column_dimensions[ws_services.cell(row=1, column=col).column_letter].width = 20
    for col in range(1, len(checklist_headers) + 1):
        ws_checklists.column_dimensions[ws_checklists.cell(row=1, column=col).column_letter].width = 25

    # Configura a resposta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="relatorio_de_servicos_Concluido_{generate_id_random()}.xlsx"'

    # Salva o arquivo Excel na resposta
    wb.save(response)

    return response