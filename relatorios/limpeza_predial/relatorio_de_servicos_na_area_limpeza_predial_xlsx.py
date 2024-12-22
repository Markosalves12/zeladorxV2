import openpyxl
from servicos.headers_report_limpeza_predial import headers_report_services
from servicos.utils_limpeza_predial import colect_dados_fato_servico_limpeza_predial
from django.http import HttpResponse
from utils.utils import generate_id_random
from datetime import datetime
from django.shortcuts import redirect
from permissionscontrol.utils import verify_login


def exportar_relatorio_de_serivos_na_area_limpeza_predial_excel(request, userid, id_random, DataDeInicio,
                                                                DataDeConclusao, Areas,
                                                                TipoServico, ServicosEscalados, ColaboradoresEscalados,
                                                                type):
    block = verify_login(request=request, userid=userid)

    if block == True:
        return redirect('logout')

    # Converte as datas de início e conclusão, tratando o caso onde a data não é fornecida
    DataDeInicio = datetime.strptime(DataDeInicio,
                                     '%Y-%m-%dT%H:%M') if DataDeInicio and DataDeInicio != "None" else None
    DataDeConclusao = datetime.strptime(DataDeConclusao,
                                        '%Y-%m-%dT%H:%M') if DataDeConclusao and DataDeConclusao != "None" else None

    # Converte os campos de Serviços e Colaboradores para listas
    ServicosEscalados = ServicosEscalados.split(',')
    ColaboradoresEscalados = ColaboradoresEscalados.split(',')

    # Coleta os dados
    dados = colect_dados_fato_servico_limpeza_predial(
        request=request,
        userid=userid,
        DataDeInicio=DataDeInicio,
        DataDeConclusao=DataDeConclusao,
        TipoServico=TipoServico,
        Areas=Areas,
        ServicosEscalados=ServicosEscalados,
        ColaboradoresEscalados=ColaboradoresEscalados,
        status=['Concluido']
    )

    wb = openpyxl.Workbook()
    ws = wb.active

    # Cabeçalhos da tabela exportada
    headers = headers_report_services

    # Preenche os cabeçalhos
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header_title

    # Aplica filtros baseados no tipo
    if type == 'catalogo_de_servicos':
        dados = dados.filter(id_random_servico=id_random)

    elif type == 'configuracao':
        dados = dados.filter(id_random_configuracao=id_random)

    elif type == 'areas':
        dados = dados.filter(id_random_area=id_random)

    # Preenche os dados no Excel
    for row_num, row in enumerate(dados, start=2):
        row_data = [
            row.tipodeempresa,
            row.empresaprestadora,
            row.id_servico,
            row.tipo_agendamento,
            row.descricao_do_servico,
            row.servicos_solicitados,
            row.data_de_inicio.replace(tzinfo=None) if row.data_de_inicio else None,  # Remove o fuso horário
            row.data_de_conclusao.replace(tzinfo=None) if row.data_de_conclusao else None,  # Remove o fuso horário
            row.status_servico,
            row.area_atendida,
            row.area_total,
            row.localidade,
            row.unidade,
            row.id_servico,
            row.tempo_na_area,
            row.colaborador_envolvido,
            row.colaborador_envolvido_id_random,
            row.foto_conclusao,
            row.data_hora_chegada.replace(tzinfo=None) if row.data_hora_chegada else None,  # Remove o fuso horário
            row.data_hora_retorno.replace(tzinfo=None) if row.data_hora_retorno else None,  # Remove o fuso horário
        ]

        # Adiciona os dados à planilha
        for col_num, value in enumerate(row_data, start=1):
            cell = ws.cell(
                row=row_num,
                column=col_num
            )
            cell.value = value

    # Configura a resposta do Excel
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response[
        'Content-Disposition'] = f'attachment; filename="relatorio_de_servicos_Concluido_{generate_id_random()}.xlsx"'

    # Salva o arquivo Excel na resposta
    wb.save(response)

    # Retorna a resposta
    return response

# def exportar_relatorio_de_serivos_na_area_limpeza_predial_excel(request, userid, id_random, DataDeInicio, DataDeConclusao, Areas,
#                                                    TipoServico, ServicosEscalados, ColaboradoresEscalados, type):
#     block = verify_login(request=request, userid=userid)
#
#     if block == True:
#         return redirect('logout')
#
#     DataDeInicio = datetime.strptime(DataDeInicio, '%Y-%m-%dT%H:%M') if DataDeInicio and DataDeInicio != "None" else 'None'
#     DataDeConclusao = datetime.strptime(DataDeConclusao, '%Y-%m-%dT%H:%M') if DataDeConclusao and DataDeConclusao != "None" else 'None'
#     ServicosEscalados = ServicosEscalados.split(',')
#     ColaboradoresEscalados = ColaboradoresEscalados.split(',')
#
#     dados = colect_dados_fato_servico_limpeza_predial(
#         request=request,
#         userid=userid,
#         DataDeInicio=DataDeInicio,
#         DataDeConclusao=DataDeConclusao,
#         TipoServico=TipoServico,
#         Areas=Areas,
#         ServicosEscalados=ServicosEscalados,
#         ColaboradoresEscalados=ColaboradoresEscalados,
#         status=['Concluido']
#     )
#
#     wb = openpyxl.Workbook()
#     ws = wb.active
#
#     # cabeçalhos da tabela exportada
#     headers = headers_report_services
#
#     for col_num, header_title in enumerate(headers, 1):
#         cell = ws.cell(row=1, column=col_num)
#         cell.value = header_title
#
#     # Adicione os dados do relatório ao arquivo Excel
#     if type == 'catalogo_de_servicos':
#         dados.filter(
#             id_random_servico=id_random
#         )
#
#     if type == 'configuracao':
#         dados.filter(
#             id_random_configuracao=id_random
#         )
#
#     elif type == 'areas':
#         dados.filter(
#             id_random_area=id_random
#         )
#
#     for row_num, row in enumerate(dados, start=2):
#         row_data = [
#             row.tipodeempresa, row.empresaprestadora,
#
#             row.id_servico, row.tipo_agendamento,
#
#             row.descricao_do_servico, row.servicos_solicitados, row.data_de_inicio,
#
#             row.data_de_conclusao, row.status_servico,
#
#             row.area_atendida, row.area_total,
#
#             row.localidade, row.unidade,
#
#             row.id_servico, row.tempo_na_area, row.colaborador_envolvido,
#
#             row.colaborador_envolvido_id_random, row.foto_conclusao,
#
#             row.data_hora_chegada, row.data_hora_retorno, row.data_hora_retorno,
#         ]
#         for col_num, value in enumerate(row_data, start=1):
#             cell = ws.cell(
#                 row=row_num,
#                 column=col_num
#             )
#             cell.value = value
#
#
#     response = HttpResponse(
#         content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
#     )
#     # cria um anexo que ´instalado no lado cliente
#     response['Content-Disposition'] = f'attachment; filename="relatorio de servicos Concluido {generate_id_random()}.xlsx"'
#
#     # Salve o arquivo Excel
#     wb.save(response)
#
#     # retorna a resposta
#     # linha obrigatória
#     # ou o sistema quebra
#     return response