import openpyxl
from servicos.headers_report_jardinagem import headers_report_services
from servicos.utils_jardinagem import colect_dados_fato_servico_jardinagem
from django.http import HttpResponse
from utils.utils import generate_id_random
from datetime import datetime


def exportar_relatorio_de_serivos_na_area_Jardinagem_excel(request, userid, id_random, DataDeInicio, DataDeConclusao, Areas,
                                                   TipoServico, ServicosEscalados, ColaboradoresEscalados, type):
    wb = openpyxl.Workbook()
    ws = wb.active

    DataDeInicio = datetime.strptime(DataDeInicio, '%Y-%m-%dT%H:%M') if DataDeInicio and DataDeInicio != "None" else 'None'
    DataDeConclusao = datetime.strptime(DataDeConclusao, '%Y-%m-%dT%H:%M') if DataDeConclusao and DataDeConclusao != "None" else 'None'
    ServicosEscalados = ServicosEscalados.split(',')
    ColaboradoresEscalados = ColaboradoresEscalados.split(',')

    filters = dict()

    if DataDeInicio and DataDeInicio != "None":
        filters['DataDeInicio__gte'] = DataDeInicio

    if DataDeConclusao and DataDeConclusao != "None":
        filters['DataDeConclusao__lte'] = DataDeConclusao

    if Areas and Areas != "None":
        filters['Areas__id__in'] = Areas

    if TipoServico and TipoServico != "None":
        filters['TipoServico'] = TipoServico

    if ServicosEscalados and ServicosEscalados != ["None"]:
        filters['ServicosEscalados__id__in'] = ServicosEscalados

    if ColaboradoresEscalados and ColaboradoresEscalados != ["None"]:
        filters['ColaboradoresEscalados__id__in'] = ColaboradoresEscalados

    dados = colect_dados_fato_servico_jardinagem(
        request=request,
        DataDeInicio=DataDeInicio,
        DataDeConclusao=DataDeConclusao,
        ServicosEscalados=ServicosEscalados,
        ColaboradoresEscalados=ColaboradoresEscalados,
        status=['Concluido'],
    )

    # cabeçalhos da tabela exportada
    headers = headers_report_services

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header_title

    # Adicione os dados do relatório ao arquivo Excel
    if type == 'catalogo_de_servicos':
        dados = dados.filter(
            id_random_servico=id_random
        )

    elif type == 'configuracao':
        dados = dados.filter(
            id_random_configuracao=id_random
        )

    elif type == 'areas':
        dados = dados.filter(
            id_random_area=id_random
        )

    elif type == 'gerente':
        dados = dados.filter(
            colaborador_envolvido_id_random=id_random
        )

    for row_num, row in enumerate(dados, start=2):
        row_data = [
            row.tipodeempresa, row.empresaprestadora,

            row.id_agendamento, row.tipo_agendamento,

            row.descricao_do_servico, row.colaboradores_chamados,

            row.servicos_solicitados, row.data_de_inicio, row.data_de_conclusao,

            row.antes, row.depois, row.status_servico, row.area_atendida, row.id_random_area,

            row.periodicidade_de_retorno, row.area_total, row.tipo_vegetacao, row.tipo_terreno,

            row.localidade, row.unidade, row.id_servico, row.tempo_na_area, row.colaborador_envolvido,

            row.colaborador_envolvido_id_random,

            row.data_hora_chegada, row.data_hora_retorno,
        ]
        for col_num, value in enumerate(row_data, start=1):
            cell = ws.cell(
                row=row_num,
                column=col_num
            )
            cell.value = value


    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    # cria um anexo que ´instalado no lado cliente
    response['Content-Disposition'] = f'attachment; filename="relatorio de servicos Concluido {generate_id_random()}.xlsx"'

    # Salve o arquivo Excel
    wb.save(response)

    # retorna a resposta
    # linha obrigatória
    # ou o sistema quebra
    return response